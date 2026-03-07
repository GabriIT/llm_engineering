#!/usr/bin/env python3
import argparse
import base64
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np
import requests
from pptx import Presentation


OLLAMA_HOST = os.environ.get("OLLAMA_HOST", "http://localhost:11434")


@dataclass
class Chunk:
    doc_id: str
    slide_no: int
    title: str
    text_extracted: str
    visual_summary: str
    key_facts: List[str]
    visual_elements: List[str]
    answerable_visual_questions: List[str]
    retrieval_text: str
    image_path: str


def slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-") or "document"


def require_binary(name: str) -> None:
    if shutil.which(name) is None:
        raise RuntimeError(
            f"Required binary '{name}' not found. Install it first, e.g. sudo apt install libreoffice poppler-utils"
        )


def run(cmd: List[str], cwd: Optional[str] = None) -> None:
    print("$", " ".join(cmd))
    subprocess.run(cmd, cwd=cwd, check=True)


def extract_slide_text(slide) -> str:
    parts: List[str] = []
    for shape in slide.shapes:
        text = getattr(shape, "text", "")
        if text and text.strip():
            cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
            if cleaned:
                parts.append(cleaned)
    return "\n\n".join(parts)


def guess_title(slide) -> str:
    title = ""
    if getattr(slide.shapes, "title", None) is not None:
        try:
            title = slide.shapes.title.text.strip()
        except Exception:
            title = ""
    if title:
        return title

    best = None
    for shape in slide.shapes:
        text = getattr(shape, "text", "").strip()
        if not text:
            continue
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        if not lines:
            continue
        candidate = lines[0]
        if best is None or len(candidate) < len(best):
            best = candidate
    return best or "Untitled Slide"


def render_slides(pptx_path: Path, output_dir: Path) -> List[Path]:
    require_binary("libreoffice")
    require_binary("pdftoppm")
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf_dir = output_dir / "pdf"
    img_dir = output_dir / "slides"
    pdf_dir.mkdir(exist_ok=True)
    img_dir.mkdir(exist_ok=True)

    run([
        "libreoffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(pdf_dir),
        str(pptx_path),
    ])

    pdf_path = pdf_dir / f"{pptx_path.stem}.pdf"
    if not pdf_path.exists():
        raise RuntimeError(f"Expected PDF not created: {pdf_path}")

    prefix = img_dir / "slide"
    run([
        "pdftoppm",
        "-png",
        "-r",
        "200",
        str(pdf_path),
        str(prefix),
    ])

    images = sorted(img_dir.glob("slide-*.png"))
    if not images:
        raise RuntimeError("No slide images rendered")
    return images


def image_to_b64(image_path: Path) -> str:
    with open(image_path, "rb") as f:
        return base64.b64encode(f.read()).decode("utf-8")


VISION_SCHEMA = {
    "type": "object",
    "properties": {
        "visual_summary": {"type": "string"},
        "key_facts": {
            "type": "array",
            "items": {"type": "string"},
        },
        "visual_elements": {
            "type": "array",
            "items": {"type": "string"},
        },
        "answerable_visual_questions": {
            "type": "array",
            "items": {"type": "string"},
        },
    },
    "required": [
        "visual_summary",
        "key_facts",
        "visual_elements",
        "answerable_visual_questions",
    ],
}


ANSWER_SCHEMA = {
    "type": "object",
    "properties": {
        "answer": {"type": "string"},
        "evidence": {
            "type": "array",
            "items": {"type": "string"},
        },
    },
    "required": ["answer", "evidence"],
}


def ollama_generate(model: str, prompt: str, images: Optional[List[Path]] = None, schema: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    payload: Dict[str, Any] = {
        "model": model,
        "prompt": prompt,
        "stream": False,
    }
    if images:
        payload["images"] = [image_to_b64(p) for p in images]
    if schema:
        payload["format"] = schema

    r = requests.post(f"{OLLAMA_HOST}/api/generate", json=payload, timeout=600)
    r.raise_for_status()
    data = r.json()
    response = data.get("response", "")
    if schema:
        try:
            return json.loads(response)
        except json.JSONDecodeError as e:
            raise RuntimeError(f"Model did not return valid JSON. Raw response:\n{response}") from e
    return {"response": response}


def ollama_embed(model: str, texts: List[str]) -> np.ndarray:
    r = requests.post(
        f"{OLLAMA_HOST}/api/embed",
        json={"model": model, "input": texts},
        timeout=600,
    )
    r.raise_for_status()
    data = r.json()
    return np.array(data["embeddings"], dtype=np.float32)


def normalize(vectors: np.ndarray) -> np.ndarray:
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    return vectors / norms


def build_chunks(pptx_path: Path, outdir: Path, vision_model: str, embed_model: str) -> None:
    prs = Presentation(str(pptx_path))
    images = render_slides(pptx_path, outdir)
    if len(images) != len(prs.slides):
        raise RuntimeError(f"Rendered {len(images)} images but PPT has {len(prs.slides)} slides")

    doc_id = slugify(pptx_path.stem)
    chunk_dir = outdir / "chunks"
    chunk_dir.mkdir(parents=True, exist_ok=True)

    chunks: List[Chunk] = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = guess_title(slide)
        text_extracted = extract_slide_text(slide)
        image_path = images[idx - 1]

        prompt = f"""
You are extracting a PowerPoint slide for multimodal RAG.

Return JSON only.
Focus on what is visually present in the slide image, especially:
- pictures, diagrams, process graphics, comparison tables, charts, labels, machine names, and numbers that are readable
- if something is uncertain, say so briefly rather than inventing

Slide title (may be incomplete): {title}
Slide text already extracted from XML:
{text_extracted[:6000]}

Create:
- visual_summary: 2-4 sentences
- key_facts: 3-8 short factual bullets
- visual_elements: objects visible on the slide, e.g. table, product photo, chart, arrow diagram
- answerable_visual_questions: 3-6 examples of questions this slide image could answer
""".strip()

        vision = ollama_generate(vision_model, prompt, images=[image_path], schema=VISION_SCHEMA)

        retrieval_text = "\n".join([
            f"title: {title}",
            f"slide_no: {idx}",
            f"text_extracted: {text_extracted}",
            f"visual_summary: {vision['visual_summary']}",
            "key_facts: " + " | ".join(vision["key_facts"]),
            "visual_elements: " + " | ".join(vision["visual_elements"]),
            "answerable_visual_questions: " + " | ".join(vision["answerable_visual_questions"]),
        ])

        chunks.append(
            Chunk(
                doc_id=doc_id,
                slide_no=idx,
                title=title,
                text_extracted=text_extracted,
                visual_summary=vision["visual_summary"],
                key_facts=vision["key_facts"],
                visual_elements=vision["visual_elements"],
                answerable_visual_questions=vision["answerable_visual_questions"],
                retrieval_text=retrieval_text,
                image_path=str(image_path),
            )
        )
        print(f"Built chunk for slide {idx}: {title}")

    jsonl_path = chunk_dir / "chunks.jsonl"
    with open(jsonl_path, "w", encoding="utf-8") as f:
        for chunk in chunks:
            f.write(json.dumps(asdict(chunk), ensure_ascii=False) + "\n")

    embeddings = normalize(ollama_embed(embed_model, [c.retrieval_text for c in chunks]))
    np.save(chunk_dir / "embeddings.npy", embeddings)
    meta = {
        "doc_id": doc_id,
        "pptx_path": str(pptx_path),
        "vision_model": vision_model,
        "embed_model": embed_model,
        "num_slides": len(chunks),
    }
    with open(chunk_dir / "meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)

    print(f"\nSaved chunks to: {jsonl_path}")
    print(f"Saved embeddings to: {chunk_dir / 'embeddings.npy'}")


def load_chunks(chunk_dir: Path) -> List[Dict[str, Any]]:
    with open(chunk_dir / "chunks.jsonl", "r", encoding="utf-8") as f:
        return [json.loads(line) for line in f]


def cosine_search(query_vec: np.ndarray, corpus: np.ndarray, top_k: int = 3) -> List[int]:
    scores = corpus @ query_vec
    idx = np.argsort(scores)[::-1][:top_k]
    return idx.tolist()


def answer_query(chunk_dir: Path, question: str, vision_model: str, embed_model: str, top_k: int = 3) -> None:
    chunks = load_chunks(chunk_dir)
    embeddings = np.load(chunk_dir / "embeddings.npy")
    q = normalize(ollama_embed(embed_model, [question]))[0]
    top_idx = cosine_search(q, embeddings, top_k=top_k)
    selected = [chunks[i] for i in top_idx]

    context_lines = []
    for ch in selected:
        context_lines.append(
            f"Slide {ch['slide_no']} | {ch['title']}\n"
            f"Text: {ch['text_extracted'][:2000]}\n"
            f"Visual summary: {ch['visual_summary']}\n"
            f"Key facts: {' | '.join(ch['key_facts'])}\n"
        )

    prompt = f"""
You are answering a question over retrieved PowerPoint slides.
Use the retrieved slide context and the retrieved slide image(s).
If the answer depends on a picture/table/graph, ground it in what is visible.
If uncertain, say what is clear and what is uncertain.
Return JSON only.

User question: {question}

Retrieved context:
{'\n---\n'.join(context_lines)}
""".strip()

    images = [Path(ch["image_path"]) for ch in selected if ch.get("image_path")]
    answer = ollama_generate(vision_model, prompt, images=images[:2], schema=ANSWER_SCHEMA)

    print("\nTop retrieved slides:")
    for ch in selected:
        print(f"- Slide {ch['slide_no']}: {ch['title']}")

    print("\nAnswer:")
    print(answer["answer"])
    print("\nEvidence:")
    for item in answer["evidence"]:
        print(f"- {item}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build and query a hybrid PPTX RAG index using Ollama vision + embeddings")
    sub = parser.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="Build chunks and embeddings from a PPTX")
    b.add_argument("--pptx", required=True, type=Path)
    b.add_argument("--outdir", required=True, type=Path)
    b.add_argument("--vision-model", default="qwen3.5:9b")
    b.add_argument("--embed-model", default="nomic-embed-text")

    q = sub.add_parser("query", help="Query a previously built chunk index")
    q.add_argument("--chunk-dir", required=True, type=Path)
    q.add_argument("--question", required=True)
    q.add_argument("--vision-model", default="qwen3.5:9b")
    q.add_argument("--embed-model", default="nomic-embed-text")
    q.add_argument("--top-k", type=int, default=3)

    args = parser.parse_args()

    if args.cmd == "build":
        build_chunks(args.pptx, args.outdir, args.vision_model, args.embed_model)
    elif args.cmd == "query":
        answer_query(args.chunk_dir, args.question, args.vision_model, args.embed_model, args.top_k)


if __name__ == "__main__":
    main()
